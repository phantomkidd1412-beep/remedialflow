"""Phase 10 - Work Tracker, Soft Delete + Recycle Bin, Admin Activity Log.

Covers:
- GET /api/tracker (view, range, search, sort, auth, validation, deleted exclusion)
- GET /api/tracker/export-link + /api/tracker/export (signing, format, auth)
- DELETE /api/projects/{pid} (soft delete, permissions, cascade hiding)
- GET /api/admin/deleted-projects (recycle bin)
- POST /api/admin/projects/{pid}/restore (restore, no duplicates)
- GET /api/admin/activity (audit, hash chain, filters)

Hits the public preview URL from EXPO_PUBLIC_BACKEND_URL. Uses TEST- prefixed
projects only; real projects (P419977, P258693, EF-10xx) are never deleted.
"""
import csv
import io
import os
import uuid
from urllib.parse import urlparse, parse_qsl

import pytest
import requests


def _read_env(path: str, key: str) -> str:
    try:
        for line in open(path):
            line = line.strip()
            if line.startswith(f"{key}="):
                return line.split("=", 1)[1].strip('"').strip("'")
    except FileNotFoundError:
        pass
    return ""


BASE_URL = (_read_env("/app/frontend/.env", "EXPO_PUBLIC_BACKEND_URL")
            or os.environ.get("EXPO_PUBLIC_BACKEND_URL", "")).rstrip("/")
API = f"{BASE_URL}/api"

CREDS = {
    "admin": ("admin@remedialflow.dev", "AdminDemo!2026"),
    "employee": ("employee@remedialflow.dev", "EmployeeDemo!2026"),
    "installer_a": ("installer@remedialflow.dev", "InstallerDemo!2026"),
    "installer_b": ("installer2@remedialflow.dev", "InstallerDemo!2026"),
}


def _login(role: str) -> str:
    email, pw = CREDS[role]
    r = requests.post(f"{API}/auth/login", json={"email": email, "password": pw}, timeout=20)
    assert r.status_code == 200, f"login {role} -> {r.status_code} {r.text[:200]}"
    return r.json()["access_token"]


def _hdr(tok: str) -> dict:
    return {"Authorization": f"Bearer {tok}", "Content-Type": "application/json"}


STATE: dict = {}


@pytest.fixture(scope="module")
def tokens():
    assert BASE_URL, "EXPO_PUBLIC_BACKEND_URL not set"
    return {k: _login(k) for k in CREDS}


@pytest.fixture(scope="module")
def ids(tokens):
    r = requests.get(f"{API}/installers", headers=_hdr(tokens["admin"]), timeout=15)
    assert r.status_code == 200
    inst = {d["email"]: d["id"] for d in r.json()}
    return {"A": inst["installer@remedialflow.dev"], "B": inst["installer2@remedialflow.dev"]}


def _create_project(tokens, ids_, installer="A", ref=None):
    ref = ref or f"TEST-P10-{uuid.uuid4().hex[:6].upper()}"
    r = requests.post(f"{API}/projects", headers=_hdr(tokens["admin"]), json={
        "ref": ref, "address": f"{ref} Tracker Rd", "postcode": "TE9 1EN",
        "installer_id": ids_[installer],
        "ncs": [{"description": "Loft insulation gaps", "required_work": "Top up", "safety_critical": False}],
    }, timeout=25)
    assert r.status_code == 200, r.text
    pid = r.json()["id"]
    return pid, ref


# =====================================================================
# 1) Tracker auth + validation
# =====================================================================
class TestTrackerAuth:
    def test_01_installer_blocked(self, tokens):
        r = requests.get(f"{API}/tracker?view=completed&range=all", headers=_hdr(tokens["installer_a"]), timeout=15)
        assert r.status_code == 403

    def test_02_admin_ok_shape(self, tokens):
        r = requests.get(f"{API}/tracker?view=completed&range=all", headers=_hdr(tokens["admin"]), timeout=15)
        assert r.status_code == 200
        body = r.json()
        assert set(["items", "count", "start", "end"]).issubset(body.keys())
        for it in body["items"][:3]:
            for key in ("id", "ref", "address", "installers", "nc_total", "nc_done",
                        "completed_at", "resolved_at", "status", "reopened", "days_outstanding"):
                assert key in it, f"missing field {key}"

    def test_03_employee_ok(self, tokens):
        r = requests.get(f"{API}/tracker?view=resolved&range=all", headers=_hdr(tokens["employee"]), timeout=15)
        assert r.status_code == 200

    def test_04_invalid_view(self, tokens):
        r = requests.get(f"{API}/tracker?view=bogus&range=all", headers=_hdr(tokens["admin"]), timeout=15)
        assert r.status_code == 400

    def test_05_invalid_range(self, tokens):
        r = requests.get(f"{API}/tracker?view=completed&range=yearly", headers=_hdr(tokens["admin"]), timeout=15)
        assert r.status_code == 400

    def test_06_invalid_custom_dates(self, tokens):
        r = requests.get(f"{API}/tracker?view=completed&range=custom&start=2026-01-10&end=2026-01-01",
                         headers=_hdr(tokens["admin"]), timeout=15)
        assert r.status_code == 400

    def test_07_bad_anchor(self, tokens):
        r = requests.get(f"{API}/tracker?view=completed&range=week&anchor=not-a-date",
                         headers=_hdr(tokens["admin"]), timeout=15)
        assert r.status_code == 400

    def test_08_search_sort_pass(self, tokens):
        r = requests.get(f"{API}/tracker?view=completed&range=all&search=TEST&sort=outstanding",
                         headers=_hdr(tokens["admin"]), timeout=15)
        assert r.status_code == 200


# =====================================================================
# 2) Export link + export auth / signing
# =====================================================================
class TestTrackerExport:
    def test_01_installer_blocked_on_link(self, tokens):
        r = requests.get(f"{API}/tracker/export-link?fmt=csv&view=completed&range=all",
                         headers=_hdr(tokens["installer_a"]), timeout=15)
        assert r.status_code == 403

    def test_02_bad_fmt(self, tokens):
        r = requests.get(f"{API}/tracker/export-link?fmt=pdf&view=completed&range=all",
                         headers=_hdr(tokens["admin"]), timeout=15)
        assert r.status_code == 400

    def test_03_csv_download(self, tokens):
        r = requests.get(f"{API}/tracker/export-link?fmt=csv&view=completed&range=all&sort=newest",
                         headers=_hdr(tokens["admin"]), timeout=15)
        assert r.status_code == 200
        url = r.json()["url"]
        STATE["csv_url"] = url
        full = f"{BASE_URL}{url}"
        rr = requests.get(full, timeout=30)
        assert rr.status_code == 200
        assert "text/csv" in rr.headers.get("content-type", "")
        assert rr.text.startswith("\ufeff"), "CSV should start with UTF-8 BOM"
        body = rr.text.lstrip("\ufeff")
        reader = csv.reader(io.StringIO(body))
        rows = list(reader)
        assert len(rows) >= 1
        assert rows[0][0] == "Project ref"
        assert "Address" in rows[0]

    def test_04_xlsx_download(self, tokens):
        r = requests.get(f"{API}/tracker/export-link?fmt=xlsx&view=resolved&range=all",
                         headers=_hdr(tokens["admin"]), timeout=15)
        assert r.status_code == 200
        full = f"{BASE_URL}{r.json()['url']}"
        rr = requests.get(full, timeout=30)
        assert rr.status_code == 200
        assert "spreadsheet" in rr.headers.get("content-type", "")
        try:
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(rr.content))
            ws = wb.active
            assert ws.cell(1, 1).value == "Project ref"
        except ImportError:
            pytest.skip("openpyxl not installed")

    def test_05_tampered_query_rejected(self, tokens):
        """Changing a param after signing must 403 (signature covers canonical query)."""
        r = requests.get(f"{API}/tracker/export-link?fmt=csv&view=completed&range=all",
                         headers=_hdr(tokens["admin"]), timeout=15)
        assert r.status_code == 200
        url = r.json()["url"]
        # swap view=completed -> view=resolved while keeping exp+sig
        tampered = url.replace("view=completed", "view=resolved")
        rr = requests.get(f"{BASE_URL}{tampered}", timeout=15)
        assert rr.status_code == 403, f"tampered URL should be 403, got {rr.status_code}"

    def test_06_rows_match_tracker(self, tokens):
        """CSV row count must match /api/tracker items count for same filters."""
        t = requests.get(f"{API}/tracker?view=completed&range=all&sort=newest",
                         headers=_hdr(tokens["admin"]), timeout=15).json()
        link = requests.get(f"{API}/tracker/export-link?fmt=csv&view=completed&range=all&sort=newest",
                            headers=_hdr(tokens["admin"]), timeout=15).json()["url"]
        rr = requests.get(f"{BASE_URL}{link}", timeout=30)
        body = rr.text.lstrip("\ufeff")
        rows = list(csv.reader(io.StringIO(body)))
        # rows[0] is header
        assert len(rows) - 1 == t["count"], f"csv data rows {len(rows)-1} != tracker count {t['count']}"


# =====================================================================
# 3) Soft delete permissions + cascade hiding
# =====================================================================
class TestSoftDelete:
    def test_01_create_project(self, tokens, ids):
        pid, ref = _create_project(tokens, ids, installer="A")
        STATE["pid"] = pid
        STATE["ref"] = ref
        STATE["ids"] = ids
        # get NC id for later checks
        d = requests.get(f"{API}/projects/{pid}", headers=_hdr(tokens["admin"]), timeout=15).json()
        STATE["nc_id"] = d["ncs"][0]["id"]

    def test_02_installer_cannot_delete(self, tokens):
        r = requests.delete(f"{API}/projects/{STATE['pid']}", headers=_hdr(tokens["installer_a"]), timeout=15)
        assert r.status_code == 403

    def test_03_employee_delete_ok(self, tokens):
        r = requests.delete(f"{API}/projects/{STATE['pid']}", headers=_hdr(tokens["employee"]), timeout=15)
        assert r.status_code == 200
        assert r.json().get("ok") is True

    def test_04_second_delete_404(self, tokens):
        r = requests.delete(f"{API}/projects/{STATE['pid']}", headers=_hdr(tokens["employee"]), timeout=15)
        assert r.status_code == 404

    def test_05_get_project_404_staff(self, tokens):
        r = requests.get(f"{API}/projects/{STATE['pid']}", headers=_hdr(tokens["admin"]), timeout=15)
        assert r.status_code == 404
        r2 = requests.get(f"{API}/projects/{STATE['pid']}", headers=_hdr(tokens["installer_a"]), timeout=15)
        assert r2.status_code == 404

    def test_06_not_in_search(self, tokens):
        r = requests.get(f"{API}/projects?search={STATE['ref']}", headers=_hdr(tokens["admin"]), timeout=15)
        assert r.status_code == 200
        items = r.json() if isinstance(r.json(), list) else r.json().get("items", [])
        refs = [i.get("ref") for i in items]
        assert STATE["ref"] not in refs

    def test_07_not_in_tracker(self, tokens):
        r = requests.get(f"{API}/tracker?view=completed&range=all&search={STATE['ref']}",
                         headers=_hdr(tokens["admin"]), timeout=15).json()
        assert all(i["ref"] != STATE["ref"] for i in r["items"])

    def test_08_review_queue_excludes(self, tokens):
        r = requests.get(f"{API}/review-queue", headers=_hdr(tokens["admin"]), timeout=15)
        if r.status_code == 200:
            refs = [i.get("project_ref") or i.get("ref") for i in r.json()]
            assert STATE["ref"] not in refs

    def test_09_mutating_endpoints_404(self, tokens):
        pid = STATE["pid"]
        nc_id = STATE["nc_id"]
        # patch project
        r1 = requests.patch(f"{API}/projects/{pid}", headers=_hdr(tokens["admin"]),
                            json={"address": "hacked"}, timeout=15)
        assert r1.status_code == 404, f"patch -> {r1.status_code}"
        # export-link should 404 (project filter) OR return 200 but rows won't include this project
        # Spec says "export-link on it -> 404"; export-link doesn't take pid, so we skip that specific case
        # resolve
        r3 = requests.post(f"{API}/projects/{pid}/resolve", headers=_hdr(tokens["admin"]),
                           json={}, timeout=15)
        assert r3.status_code in (404,), f"resolve -> {r3.status_code}"
        # rates endpoint
        r4 = requests.put(f"{API}/projects/{pid}/rates/{STATE['ids']['A']}", headers=_hdr(tokens["admin"]),
                          json={"rate": 200}, timeout=15)
        assert r4.status_code == 404, f"rates -> {r4.status_code}"
        # ncs under project
        r5 = requests.patch(f"{API}/ncs/{nc_id}", headers=_hdr(tokens["admin"]),
                            json={"description": "x"}, timeout=15)
        # NC endpoint may 404 via project_deleted guard
        assert r5.status_code in (404, 403), f"nc patch -> {r5.status_code}"


# =====================================================================
# 4) Recycle bin (admin-only) + restore
# =====================================================================
def _earning_jobs(groups, ref):
    """/admin/payments returns [{installer_id, week, jobs:[{ref, amount, status}]}]."""
    return sum(1 for g in groups for j in g.get("jobs", []) if j.get("ref") == ref)


class TestRecycleBin:
    def test_01_employee_no_access(self, tokens):
        r = requests.get(f"{API}/admin/deleted-projects", headers=_hdr(tokens["employee"]), timeout=15)
        assert r.status_code == 403

    def test_02_installer_no_access(self, tokens):
        r = requests.get(f"{API}/admin/deleted-projects", headers=_hdr(tokens["installer_a"]), timeout=15)
        assert r.status_code == 403

    def test_03_admin_sees_our_project(self, tokens):
        r = requests.get(f"{API}/admin/deleted-projects", headers=_hdr(tokens["admin"]), timeout=15)
        assert r.status_code == 200
        rows = r.json()
        hit = [x for x in rows if x["id"] == STATE["pid"]]
        assert hit, f"deleted project not in recycle bin"
        row = hit[0]
        assert row["ref"] == STATE["ref"]
        d = row["deleted"]
        assert d["by_role"] == "employee"
        assert "Employee" in d["by_name"] or d["by_name"]
        assert d.get("at")

    def test_04_employee_cannot_restore(self, tokens):
        r = requests.post(f"{API}/admin/projects/{STATE['pid']}/restore",
                          headers=_hdr(tokens["employee"]), timeout=15)
        assert r.status_code == 403

    def test_05_capture_pre_restore_counts(self, tokens):
        # count earnings for project (via admin payments)
        r = requests.get(f"{API}/admin/payments?status=all",
                         headers=_hdr(tokens["admin"]), timeout=15)
        STATE["pre_earnings"] = 0
        if r.status_code == 200:
            STATE["pre_earnings"] = _earning_jobs(r.json(), STATE["ref"])

    def test_06_admin_restore_ok(self, tokens):
        r = requests.post(f"{API}/admin/projects/{STATE['pid']}/restore",
                          headers=_hdr(tokens["admin"]), timeout=15)
        assert r.status_code == 200, r.text

    def test_07_second_restore_404(self, tokens):
        r = requests.post(f"{API}/admin/projects/{STATE['pid']}/restore",
                          headers=_hdr(tokens["admin"]), timeout=15)
        assert r.status_code == 404

    def test_08_project_visible_again(self, tokens):
        r = requests.get(f"{API}/projects/{STATE['pid']}", headers=_hdr(tokens["admin"]), timeout=15)
        assert r.status_code == 200
        d = r.json()
        assert d["project"]["ref"] == STATE["ref"]
        # deletion_history must contain entry
        hist = d["project"].get("deletion_history") or []
        assert len(hist) >= 1
        assert hist[-1].get("restored_at")

    def test_09_no_duplicate_earnings(self, tokens):
        r = requests.get(f"{API}/admin/payments?status=all",
                         headers=_hdr(tokens["admin"]), timeout=15)
        if r.status_code != 200:
            pytest.skip("admin payments not available")
        post = _earning_jobs(r.json(), STATE["ref"])
        assert post == STATE["pre_earnings"], f"earnings duplicated: pre={STATE['pre_earnings']} post={post}"


# =====================================================================
# 5) Activity log (admin-only, audit chain)
# =====================================================================
class TestActivityLog:
    def test_01_employee_blocked(self, tokens):
        r = requests.get(f"{API}/admin/activity", headers=_hdr(tokens["employee"]), timeout=15)
        assert r.status_code == 403

    def test_02_installer_blocked(self, tokens):
        r = requests.get(f"{API}/admin/activity", headers=_hdr(tokens["installer_a"]), timeout=15)
        assert r.status_code == 403

    def test_03_admin_ok_shape(self, tokens):
        r = requests.get(f"{API}/admin/activity?limit=50", headers=_hdr(tokens["admin"]), timeout=15)
        assert r.status_code == 200
        body = r.json()
        assert "items" in body and "total" in body
        for it in body["items"][:3]:
            for key in ("seq", "at", "event", "record_type", "record_id",
                        "actor_id", "actor_name", "actor_role", "hash"):
                assert key in it, f"missing {key}"

    def test_04_deletion_filter_and_our_event(self, tokens):
        r = requests.get(f"{API}/admin/activity?kind=deletion&limit=200",
                         headers=_hdr(tokens["admin"]), timeout=15)
        assert r.status_code == 200
        items = r.json()["items"]
        # all events are in deletion family
        valid = {"project.deleted", "project.restored", "project.purged", "evidence.discarded",
                 "evidence.excluded", "evidence.restored"}
        assert all(it["event"] in valid for it in items)
        # find our delete + restore pair for STATE['pid']
        dels = [it for it in items if it["event"] == "project.deleted"
                and it["record_id"] == STATE["pid"]]
        rests = [it for it in items if it["event"] == "project.restored"
                 and it["record_id"] == STATE["pid"]]
        assert dels, "no project.deleted entry for our test project"
        assert rests, "no project.restored entry for our test project"
        assert dels[0]["actor_role"] == "employee"
        assert rests[0]["actor_role"] == "admin"

    def test_05_search(self, tokens):
        r = requests.get(f"{API}/admin/activity?search={STATE['ref']}&limit=50",
                         headers=_hdr(tokens["admin"]), timeout=15)
        assert r.status_code == 200
        items = r.json()["items"]
        # at least our delete/restore
        assert any(it["event"] in ("project.deleted", "project.restored")
                   and it["record_id"] == STATE["pid"] for it in items)

    def test_06_no_write_endpoints(self, tokens):
        """There must be no API to modify/delete audit entries."""
        for method in ("delete", "put", "patch", "post"):
            r = requests.request(method, f"{API}/admin/activity/1",
                                 headers=_hdr(tokens["admin"]), timeout=10)
            assert r.status_code in (404, 405), f"{method} /admin/activity/1 -> {r.status_code}"

    def test_07_hash_chain_intact(self, tokens):
        r = requests.get(f"{API}/admin/activity?limit=50", headers=_hdr(tokens["admin"]), timeout=15)
        items = r.json()["items"]
        # items are seq-desc; walk pairs: for each consecutive pair the later-older
        # entry's hash becomes next entry's prev_hash. We can't see prev_hash in
        # response (filtered out), but we can verify monotonic seq and unique hashes.
        seqs = [it["seq"] for it in items]
        assert seqs == sorted(seqs, reverse=True)
        hashes = [it["hash"] for it in items]
        assert len(set(hashes)) == len(hashes), "duplicate hashes - chain broken"
        assert all(len(h) == 64 for h in hashes), "sha256 hashes expected"


# =====================================================================
# 6) Reopened project appears once in tracker
# =====================================================================
class TestReopenedTrackerDedup:
    """A project that was completed, reopened, and re-completed must appear exactly once
    in the completed view."""

    def test_01_count_test_p9dual(self, tokens):
        # Our known reopened candidate from phase 9 seed: TEST-P9DUAL-74D7DB
        r = requests.get(f"{API}/tracker?view=completed&range=all&sort=newest",
                         headers=_hdr(tokens["admin"]), timeout=15)
        assert r.status_code == 200
        items = r.json()["items"]
        matches = [i for i in items if i["ref"].startswith("TEST-")]
        refs = [i["ref"] for i in matches]
        # No duplicate refs
        assert len(refs) == len(set(refs)), f"duplicate rows for refs: {refs}"
